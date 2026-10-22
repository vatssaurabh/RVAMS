"""
export_data.py — Step 7 & Step 8: Export verified student list, long-format attendance,
register-style attendance matrix, review report, and multi-sheet Excel workbook.
"""

from __future__ import annotations

import io
import os
from dataclasses import dataclass
from typing import Dict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd


@dataclass
class ExportedFilesResult:
    output_dir: str
    file_paths: Dict[str, str]
    excel_bytes: bytes


def _format_excel_workbook(wb: openpyxl.Workbook) -> None:
    """Apply clean header formatting, column widths, and text-format Roll_No cells."""
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    present_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    absent_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue

        roll_col_indices = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if str(cell.value) == "Roll_No":
                roll_col_indices.append(col_idx)

        for row_idx in range(2, ws.max_row + 1):
            for col_idx in roll_col_indices:
                rcell = ws.cell(row=row_idx, column=col_idx)
                rcell.value = str(rcell.value) if rcell.value is not None else ""
                rcell.number_format = "@"

            for col_idx in range(1, ws.max_column + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                val_str = str(c.value or "").strip()
                if val_str == "Present":
                    c.fill = present_fill
                elif val_str == "Absent":
                    c.fill = absent_fill

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_len = len(str(cell.value or ""))
                if val_len > max_len:
                    max_len = val_len
            ws.column_dimensions[col_letter].width = max(12, min(35, max_len + 4))


def export_all_reports(
    students_df: pd.DataFrame,
    attendance_long_df: pd.DataFrame,
    attendance_matrix_df: pd.DataFrame,
    review_report_df: pd.DataFrame,
    output_dir: str = "output",
) -> ExportedFilesResult:
    """
    Save all 5 expected output files specified in Section 6 of the project guide:
      - `students.csv`
      - `attendance.csv`
      - `attendance_matrix.csv`
      - `review_report.csv`
      - `attendance.xlsx`
    """
    os.makedirs(output_dir, exist_ok=True)

    # Ensure Roll_No is string across all DataFrames
    s_df = students_df.copy()
    a_df = attendance_long_df.copy()
    m_df = attendance_matrix_df.copy()
    r_df = review_report_df.copy()

    for df in (s_df, a_df, m_df, r_df):
        if "Roll_No" in df.columns:
            df["Roll_No"] = df["Roll_No"].astype(str)

    students_csv_path = os.path.join(output_dir, "students.csv")
    attendance_csv_path = os.path.join(output_dir, "attendance.csv")
    matrix_csv_path = os.path.join(output_dir, "attendance_matrix.csv")
    review_csv_path = os.path.join(output_dir, "review_report.csv")
    excel_path = os.path.join(output_dir, "attendance.xlsx")

    s_df.to_csv(students_csv_path, index=False, encoding="utf-8-sig")
    a_df.to_csv(attendance_csv_path, index=False, encoding="utf-8-sig")
    m_df.to_csv(matrix_csv_path, index=False, encoding="utf-8-sig")
    r_df.to_csv(review_csv_path, index=False, encoding="utf-8-sig")

    excel_buffer = io.BytesIO()
    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        m_df.to_excel(writer, sheet_name="Attendance_Matrix", index=False)
        s_df.to_excel(writer, sheet_name="Students", index=False)
        a_df.to_excel(writer, sheet_name="Attendance_Long", index=False)
        r_df.to_excel(writer, sheet_name="Review_Report", index=False)
        _format_excel_workbook(writer.book)

    excel_bytes = excel_buffer.getvalue()
    with open(excel_path, "wb") as f:
        f.write(excel_bytes)

    return ExportedFilesResult(
        output_dir=output_dir,
        file_paths={
            "students.csv": students_csv_path,
            "attendance.csv": attendance_csv_path,
            "attendance_matrix.csv": matrix_csv_path,
            "review_report.csv": review_csv_path,
            "attendance.xlsx": excel_path,
        },
        excel_bytes=excel_bytes,
    )
